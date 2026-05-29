"""
Multi-modal document extraction (AWA-6).

``INGEST_EXTRACTOR``:
  - ``native`` (default): HTML / plain text / PDF via PyMuPDF + Tesseract fallback.
  - ``gemini``: legacy per-page PDF extraction via Gemini Flash (see ``extractor.py``).
"""

from __future__ import annotations

import asyncio
import logging
import os
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from xml.etree import ElementTree as ET

import fitz  # PyMuPDF
from bs4 import BeautifulSoup
from google import genai

from ai_engine.extractor import PageText
from ai_engine.extractor import extract_text as extract_text_gemini
from ai_engine.text_quality import assess_extracted_text, is_corrupt_extracted_text
from ai_engine.text_utils import normalize_unicode_nfc

logger = logging.getLogger(__name__)

INGEST_EXTRACTOR = os.getenv("INGEST_EXTRACTOR", "native").lower()
MIN_TEXT_CHARS_FOR_NATIVE_PAGE = int(os.getenv("MIN_TEXT_CHARS_PER_PAGE", "40"))
OCR_LANGS = os.getenv("OCR_LANGS", "eng+amh")
OCR_CONFIDENCE_FAIL_THRESHOLD = float(os.getenv("OCR_CONFIDENCE_FAIL_THRESHOLD", "0.5"))
OCR_RENDER_SCALE = float(os.getenv("OCR_RENDER_SCALE", "3.0"))
# --oem 1  → LSTM neural net engine only (best accuracy, trained on real text)
# --psm 3  → Fully automatic page segmentation (handles mixed Amharic+English columns)
OCR_CONFIG = os.getenv("OCR_CONFIG", "--oem 1 --psm 3")

_tesseract_langs_cache: set[str] | None = None
_amh_warning_logged = False


@dataclass
class ExtractionOutcome:
    pages: list[PageText]
    requires_manual_review: bool
    review_reason: str | None = None


def _mean(values: list[float]) -> float | None:
    if not values:
        return None
    return sum(values) / len(values)


def get_tesseract_languages() -> set[str]:
    """Installed Tesseract language codes (cached)."""
    global _tesseract_langs_cache
    if _tesseract_langs_cache is not None:
        return _tesseract_langs_cache
    try:
        import pytesseract

        _tesseract_langs_cache = set(pytesseract.get_languages(config=""))
    except Exception:
        logger.exception("Failed to list Tesseract languages")
        _tesseract_langs_cache = set()
    return _tesseract_langs_cache


def resolve_ocr_langs(requested: str | None = None) -> str:
    """Map ``OCR_LANGS`` to installed packs; warn when ``amh`` is missing."""
    global _amh_warning_logged
    raw = (requested or OCR_LANGS).strip()
    parts = [p.strip() for p in raw.split("+") if p.strip()]
    available = get_tesseract_languages()
    resolved = [p for p in parts if p in available] or ["eng"]
    if "amh" in parts and "amh" not in available and not _amh_warning_logged:
        _amh_warning_logged = True
        logger.warning(
            "Tesseract 'amh' traineddata is not installed (have: %s). "
            "Amharic PDFs will OCR poorly. Install e.g. "
            "'brew install tesseract-lang' or 'apt install tesseract-ocr-amh'.",
            sorted(available),
        )
    return "+".join(resolved)


def _preprocess_for_ocr(img: "Image.Image") -> "Image.Image":
    """
    Improve OCR accuracy on scanned documents via:
      1. Grayscale — removes colour noise
      2. Gaussian blur — reduces scan grain before thresholding
      3. Otsu binarization — converts to pure black/white, maximises contrast
      4. Upscale if small — Tesseract works best at ≥300 DPI equivalent
    Especially important for Amharic: the `amh` traineddata was trained on
    clean binarised text; feeding it a noisy greyscale scan hurts accuracy.
    """
    import numpy as np
    from PIL import Image

    # Convert to grayscale
    gray = img.convert("L")

    # Upscale small images — Tesseract needs enough pixels per character
    w, h = gray.size
    if w < 1800:
        scale = max(2, 1800 // w)
        gray = gray.resize((w * scale, h * scale), Image.LANCZOS)

    # Gaussian blur to reduce noise before threshold
    from PIL import ImageFilter
    blurred = gray.filter(ImageFilter.GaussianBlur(radius=1))

    # Otsu binarisation via numpy
    arr = np.array(blurred, dtype=np.uint8)
    hist, _ = np.histogram(arr.flatten(), bins=256, range=(0, 256))
    total = arr.size
    best_thresh, best_var = 0, 0.0
    w_bg = 0.0
    sum_total = float(np.dot(np.arange(256), hist))
    sum_bg = 0.0
    for t in range(256):
        w_bg += hist[t]
        if w_bg == 0:
            continue
        w_fg = total - w_bg
        if w_fg == 0:
            break
        sum_bg += t * hist[t]
        mean_bg = sum_bg / w_bg
        mean_fg = (sum_total - sum_bg) / w_fg
        var = (w_bg / total) * (w_fg / total) * (mean_bg - mean_fg) ** 2
        if var > best_var:
            best_var, best_thresh = var, t
    binarised = arr > best_thresh
    result = Image.fromarray((binarised * 255).astype(np.uint8))
    return result


def ocr_pil_image(img: "Image.Image") -> tuple[str, float]:
    """Run Tesseract on a PIL image; return (text, mean_confidence 0..1)."""
    import pytesseract
    from PIL import Image

    img = _preprocess_for_ocr(img)
    lang = resolve_ocr_langs()
    data = pytesseract.image_to_data(img, lang=lang, config=OCR_CONFIG, output_type=pytesseract.Output.DICT)
    confs: list[float] = []
    words: list[str] = []
    for i, conf in enumerate(data.get("conf", [])):
        try:
            c = float(conf)
        except (TypeError, ValueError):
            continue
        if c < 0:
            continue
        confs.append(c / 100.0)
        w = (data.get("text") or [""])[i]
        if w and w.strip():
            words.append(w)
    text = " ".join(words).strip()
    mean_conf = float(sum(confs) / len(confs)) if confs else 0.0
    return text, mean_conf


def ocr_image_bytes(data: bytes) -> tuple[str, float]:
    """OCR raw image bytes (JPEG/PNG/WebP)."""
    from PIL import Image

    img = Image.open(BytesIO(data))
    return ocr_pil_image(img)


def _ocr_page(pdf_page: fitz.Page) -> tuple[str, float]:
    """Render page and OCR with Tesseract; return (text, mean_confidence 0..1)."""
    from PIL import Image

    scale = max(1.5, OCR_RENDER_SCALE)
    mat = fitz.Matrix(scale, scale)
    pix = pdf_page.get_pixmap(matrix=mat, alpha=False)
    mode = "RGB" if pix.n == 3 else "RGBA"
    img = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
    if mode == "RGBA":
        img = img.convert("RGB")

    return ocr_pil_image(img)


def _extract_pdf_native(pdf_bytes: bytes) -> ExtractionOutcome:
    """PyMuPDF text layer with OCR fallback per page."""
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages: list[PageText] = []
    ocr_confidences: list[float] = []

    used_corrupt_embedded = False
    used_poor_ocr = False

    try:
        for i in range(len(doc)):
            page = doc[i]
            raw = normalize_unicode_nfc((page.get_text("text") or "").strip())
            use_native = len(raw) >= MIN_TEXT_CHARS_FOR_NATIVE_PAGE
            if use_native and is_corrupt_extracted_text(raw):
                used_corrupt_embedded = True
                logger.info(
                    "Page %d: embedded PDF text looks corrupt; falling back to OCR",
                    i + 1,
                )
                use_native = False

            if use_native:
                pages.append(
                    PageText(
                        page_number=i + 1,
                        text=raw,
                        extraction_mode="pdf_text",
                        ocr_mean_confidence=None,
                    )
                )
                continue

            ocr_text, ocr_mean = _ocr_page(page)
            ocr_text = normalize_unicode_nfc(ocr_text)
            ocr_confidences.append(ocr_mean)
            if is_corrupt_extracted_text(ocr_text):
                used_poor_ocr = True
            pages.append(
                PageText(
                    page_number=i + 1,
                    text=ocr_text,
                    extraction_mode="ocr",
                    ocr_mean_confidence=ocr_mean,
                )
            )
    finally:
        doc.close()

    if not any(p.text.strip() for p in pages):
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )

    full_doc_text = "\n".join(p.text for p in pages if p.text.strip())
    doc_quality = assess_extracted_text(full_doc_text)

    review_reason: str | None = None
    mean_ocr = _mean(ocr_confidences)

    if doc_quality.is_likely_corrupt or used_poor_ocr:
        review_reason = "poor_text_quality"
    elif used_corrupt_embedded:
        review_reason = "corrupt_embedded_text"
    elif mean_ocr is not None and mean_ocr < OCR_CONFIDENCE_FAIL_THRESHOLD:
        review_reason = "low_ocr_confidence"
    elif "amh" not in get_tesseract_languages() and doc_quality.ethiopic_ratio < 0.05:
        review_reason = "amh_traineddata_missing"

    if review_reason:
        return ExtractionOutcome(
            pages=pages,
            requires_manual_review=True,
            review_reason=review_reason,
        )

    return ExtractionOutcome(pages=pages, requires_manual_review=False)


def _extract_html(html_bytes: bytes) -> ExtractionOutcome:
    soup = BeautifulSoup(html_bytes, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    text = normalize_unicode_nfc(re.sub(r"\n{3,}", "\n\n", text))
    pages = [
        PageText(
            page_number=1,
            text=text,
            extraction_mode="html",
            ocr_mean_confidence=None,
        )
    ]
    return ExtractionOutcome(pages=pages, requires_manual_review=False)


def _extract_plain_text(data: bytes) -> ExtractionOutcome:
    text = normalize_unicode_nfc(data.decode("utf-8", errors="replace"))
    pages = [
        PageText(
            page_number=1,
            text=text,
            extraction_mode="manual_txt",
            ocr_mean_confidence=None,
        )
    ]
    return ExtractionOutcome(pages=pages, requires_manual_review=False)


def _extract_image(data: bytes) -> ExtractionOutcome:
    """OCR a single image (Telegram photos, scanned pages as PNG/JPEG)."""
    try:
        ocr_text, ocr_mean = ocr_image_bytes(data)
    except Exception:
        logger.exception("image_ocr_failed")
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )
    ocr_text = normalize_unicode_nfc(ocr_text)
    if not ocr_text.strip():
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )
    review_reason: str | None = None
    if is_corrupt_extracted_text(ocr_text):
        review_reason = "poor_text_quality"
    elif ocr_mean < OCR_CONFIDENCE_FAIL_THRESHOLD:
        review_reason = "low_ocr_confidence"
    pages = [
        PageText(
            page_number=1,
            text=ocr_text,
            extraction_mode="ocr",
            ocr_mean_confidence=ocr_mean,
        )
    ]
    if review_reason:
        return ExtractionOutcome(
            pages=pages,
            requires_manual_review=True,
            review_reason=review_reason,
        )
    return ExtractionOutcome(pages=pages, requires_manual_review=False)


def _extract_pptx(data: bytes) -> ExtractionOutcome:
    """Extract slide text from OOXML presentations."""
    try:
        from pptx import Presentation
    except ImportError as e:
        raise RuntimeError("python-pptx is required for PPTX ingestion") from e

    try:
        text_parts: list[str] = []
        prs = Presentation(BytesIO(data))
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_lines: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text and text.strip():
                    slide_lines.append(text.strip())
            if slide_lines:
                text_parts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_lines))
        text = normalize_unicode_nfc("\n\n".join(text_parts))
    except (KeyError, OSError, ValueError, AttributeError):
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )
    if not text.strip():
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )
    pages = [
        PageText(
            page_number=1,
            text=text,
            extraction_mode="pptx",
            ocr_mean_confidence=None,
        )
    ]
    return ExtractionOutcome(pages=pages, requires_manual_review=False)


def _extract_docx(data: bytes) -> ExtractionOutcome:
    """Pull paragraph text from OOXML (no external deps)."""
    try:
        text_parts: list[str] = []
        with zipfile.ZipFile(BytesIO(data)) as zf:
            with zf.open("word/document.xml") as f:
                tree = ET.parse(f)
                root = tree.getroot()
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                for node in root.findall(".//w:p", ns):
                    runs = []
                    for t in node.findall(".//w:t", ns):
                        if t.text:
                            runs.append(t.text)
                    if runs:
                        text_parts.append("".join(runs))
        text = normalize_unicode_nfc("\n".join(text_parts))
    except (KeyError, OSError, ET.ParseError, zipfile.BadZipFile):
        return ExtractionOutcome(
            pages=[],
            requires_manual_review=False,
            review_reason="empty_extract",
        )
    pages = [
        PageText(
            page_number=1,
            text=text,
            extraction_mode="docx",
            ocr_mean_confidence=None,
        )
    ]
    return ExtractionOutcome(pages=pages, requires_manual_review=False)


async def extract_bytes(
    data: bytes,
    *,
    mime_type: str | None,
    filename: str | None,
    genai_client: genai.Client | None = None,
) -> ExtractionOutcome:
    """
    Route extraction based on MIME / filename.

    PDF:
      - ``INGEST_EXTRACTOR=gemini`` → Gemini per-page extraction.
      - otherwise → native PyMuPDF + OCR fallback.
    """
    name = (filename or "").lower()
    mt = (mime_type or "").split(";")[0].strip().lower()

    if mt in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ) or name.endswith(".docx"):
        return await asyncio.to_thread(_extract_docx, data)

    if mt in (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
    ) or name.endswith(".pptx"):
        return await asyncio.to_thread(_extract_pptx, data)

    if mt.startswith("image/") or name.endswith((".jpg", ".jpeg", ".png", ".webp")):
        return await asyncio.to_thread(_extract_image, data)

    if mt == "application/pdf" or name.endswith(".pdf"):
        if INGEST_EXTRACTOR == "gemini":
            client = genai_client or genai.Client()
            pages = await extract_text_gemini(data, client=client)
            for p in pages:
                if not p.extraction_mode:
                    p.extraction_mode = "gemini"
            return ExtractionOutcome(pages=pages, requires_manual_review=False)
        return await asyncio.to_thread(_extract_pdf_native, data)

    if mt in ("text/html", "application/xhtml+xml") or name.endswith((".html", ".htm")):
        return await asyncio.to_thread(_extract_html, data)

    if mt.startswith("text/") or name.endswith(".txt"):
        return await asyncio.to_thread(_extract_plain_text, data)

    # Default: try UTF-8 text
    return await asyncio.to_thread(_extract_plain_text, data)


class DocumentProcessor:
    """SDS-aligned façade for multi-modal extraction (same behavior as ``extract_bytes``)."""

    @staticmethod
    async def extract_bytes(
        data: bytes,
        *,
        mime_type: str | None,
        filename: str | None,
        genai_client: genai.Client | None = None,
    ) -> ExtractionOutcome:
        return await extract_bytes(data, mime_type=mime_type, filename=filename, genai_client=genai_client)
